"""
Export Excel multi-onglets (§20) : Ranking, Technical Details, Fundamentals
Details, Catalyst Details, Execution Plan, Data Quality (+ onglet Params).
"""
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

DECISION_COLORS = {
    "Achat fort": "C6EFCE",
    "Achat progressif": "DDEBF7",
    "Watchlist prioritaire": "FFF2CC",
    "Surveillance simple": "F2F2F2",
    "Éviter": "F8CBAD",
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_sheet(ws, df: pd.DataFrame, color_col: str | None = None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        content_len = df[col].astype(str).str.len().quantile(0.9) if len(df) else 0
        if pd.isna(content_len):
            content_len = 0
        width = max(len(str(col)) + 2, min(int(content_len) + 2, 55))
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    if color_col and color_col in df.columns:
        j = list(df.columns).index(color_col) + 1
        for i, val in enumerate(df[color_col], start=2):
            color = DECISION_COLORS.get(val)
            if color:
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=color)


def _fmt_pct(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    df = df.copy()
    for c in cols:
        if c in df.columns:
            df[c] = df[c].map(lambda v: round(v * 100, 2) if pd.notna(v) else None)
            df.rename(columns={c: c + "_%"}, inplace=True)
    return df


def export(master: pd.DataFrame, quality: pd.DataFrame, meta: dict,
           path: str = config.OUTPUT_FILE) -> str:
    m = master.reset_index().rename(columns={"index": "ticker"})

    ranking = m[["rank", "ticker", "company_name", "sector", "liquidity_class",
                 "halal_status", "halal_alert", "stage_weinstein", "graham_score",
                 "technical_score", "catalyst_score", "liquidity_execution_score",
                 "final_score", "halal_adjusted_score", "decision",
                 "max_position_size", "execution_note", "sell_signal",
                 "short_comment"]].sort_values("rank")

    technical = _fmt_pct(
        m[["ticker", "close", "ma30w", "ma30w_slope", "price_above_ma30",
           "distance_to_ma30", "rs_26w_vs_brvm30", "stage_weinstein",
           "volume_confirmation_score", "technical_reliability_score",
           "technical_score_raw", "technical_score"]],
        ["distance_to_ma30", "rs_26w_vs_brvm30"])

    fundamentals = _fmt_pct(
        m[["ticker", "sector", "is_bank", "per", "pb", "roe", "roa",
           "dividend_yield", "payout_ratio", "net_debt_ebitda", "debt_to_equity",
           "eps_growth", "graham_coverage", "graham_score", "graham_comment"]],
        ["roe", "roa", "dividend_yield", "payout_ratio", "eps_growth"])

    catalyst = _fmt_pct(
        m[["ticker", "net_income_yoy_growth", "expected_dividend_change",
           "implied_dividend_yield", "next_event_date", "event_type",
           "market_reaction_score", "catalyst_coverage", "catalyst_score",
           "catalyst_comment"]],
        ["net_income_yoy_growth", "expected_dividend_change", "implied_dividend_yield"])

    execution = m[["ticker", "liquidity_class", "trading_frequency",
                   "avg_weekly_value", "liquidity_execution_score",
                   "max_position_size", "suggested_entry_tranches",
                   "order_type", "execution_note", "risk_note"]]

    params_df = pd.DataFrame(
        [{"paramètre": k, "valeur": v} for k, v in config.PARAMS.items()]
        + [{"paramètre": f"meta_{k}", "valeur": str(v)} for k, v in meta.items()])

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        sheets = [("01_Ranking", ranking, "decision"),
                  ("02_Technical_Details", technical, None),
                  ("03_Fundamentals_Details", fundamentals, None),
                  ("04_Catalyst_Details", catalyst, None),
                  ("05_Execution_Plan", execution, None),
                  ("06_Data_Quality", quality.reset_index(), None),
                  ("07_Params", params_df, None)]
        for name, df, color_col in sheets:
            df.to_excel(xl, sheet_name=name, index=False)
            _style_sheet(xl.sheets[name], df, color_col)
    return path
